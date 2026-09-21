from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

base = Path(__file__).parent

# Read memoria text
memoria_txt = base / 'p1_memoria.txt'
if memoria_txt.exists():
    text = memoria_txt.read_text(encoding='utf-8')
else:
    text = 'Memoria no encontrada.'

# Create DOCX
docx_path = base / 'dawm2d_opia_p1.docx'
doc = Document()
doc.add_heading('PRACTICA 1 - Memoria', level=1)
for line in text.splitlines():
    doc.add_paragraph(line)
doc.add_page_break()
doc.save(docx_path)

# Create simple PDF
pdf_path = base / 'dawm2d_opia_p1.pdf'
c = canvas.Canvas(str(pdf_path), pagesize=A4)
width, height = A4
y = height - 72
for line in text.splitlines():
    c.drawString(72, y, line)
    y -= 14
    if y < 72:
        c.showPage()
        y = height - 72
c.save()

print('Created', docx_path.name, pdf_path.name)
